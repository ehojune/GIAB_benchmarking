#!/usr/bin/env python3
"""phase3 samplesheets -> 파이프라인 입력 디렉토리, 샘플당 FASTQ 한 쌍 (nbb2에서 실행).

    cd <파이프라인 작업 디렉토리>            # 예: G000-gd1-20260819 옆
    python <repo>/phase3_shortread_wgs/scripts/03_make_pipeline_dirs.py --dry-run      # 계획만
    python .../03_make_pipeline_dirs.py --prune                                        # 1쌍짜리는 링크, 여러 쌍은 concat.sh 작성 (옛 링크 정리)
    python .../03_make_pipeline_dirs.py --prune --run-concat --jobs 4                  # + 병합을 여기서 4개씩 병렬로 전부 실행하고 기다린다 (nohup 권장)
    python .../03_make_pipeline_dirs.py --prune --qsub                                 # + 병합을 샘플별 SGE 잡으로 제출 (shepherd.q)
    python .../03_make_pipeline_dirs.py --only Hiseq-300x --only Illumina-250PE        # 일부 데이터셋만(set 이름 뒷부분)
    python .../03_make_pipeline_dirs.py --prefix GIAB-publicdata                       # 접두 변경 (기본 GIAB-publicData)

만드는 구조 (업체 디렉토리 G000-gd?-*/outcome/<sample>/<sample>_1.fastq.gz 와 같은 모양, **샘플당 R1/R2 한 쌍**).
이름 안의 구분자는 전부 `-`이고 `_`는 mate 접미(`_1.fastq.gz`/`_2.fastq.gz`)에만 쓴다(2026-09-19 사용자 결정 — 파이프라인이 `_`로 샘플명을 자른다):

    <dest>/<prefix>-<dataset>/outcome/<prefix>-<dataset>-<HG00?>/<prefix>-<dataset>-<HG00?>_1.fastq.gz
                                                                                  ..._2.fastq.gz
    예: GIAB-publicData-Novaseq6000-PCRfree-30x/outcome/GIAB-publicData-Novaseq6000-PCRfree-30x-HG002/GIAB-publicData-Novaseq6000-PCRfree-30x-HG002_1.fastq.gz

- 샘플에 쌍이 하나면(NovaSeq 6000/X, HiSeq 30x 서브샘플) 원본으로의 심볼릭 링크(`--hard`면 하드링크).
- 쌍이 여럿이면(HiSeq300x 935~1,024쌍, 2x250 18~35쌍, MGISEQ·BGISEQ·AVITI 2~4쌍) **cat으로 한 쌍에 병합**한다(2026-09-18 사용자 결정 "각각 알아서").
  샘플 dir에 concat_R1.list / concat_R2.list(같은 순서: unit, 경로)와 concat.sh를 쓴다. gzip은 멀티멤버라 이어붙여도 유효하다.
  concat.sh는 출력 크기 = 입력 합을 확인하고, 이미 맞으면 SKIP한다(재실행 안전). 병합 파일은 실사본이다(HiSeq300x 3샘플만 2.4 TiB).
  AVITI는 StdInsert(~400 bp)와 LngInsert(~1300 bp) 라이브러리가 한 파일로 합쳐진다 — 인서트 분포가 이봉이 된다는 점을 파이프라인 해석에 남길 것.
- 원본은 samplesheets/<dsid>.csv 의 절대경로 그대로다(DATA_ROOT=/BiO/scratch/ehojune/GIAB_benchmark).
- 멱등: 같은 대상을 가리키는 링크는 건너뛰고, 다른 대상을 가리키면 멈춘다. `--prune`은 관리하는 샘플 dir 안의 계획에 없는 심볼릭 링크만 지운다
  (병합 결과·일반 파일은 안 건드림). 이름 규칙이 바뀐 뒤 재실행할 때 쓴다.

2차(2026-09-23, HG001·HG005~HG009 + HG002 잔여): more_run_table.tsv(make_more_samplesheets.py 생성)에 있는 dsid는 DIRNAME 대신
그 표의 set/label로 간다 — `<prefix>-<set>/outcome/<prefix>-<set>-<label>/…`. 예: GIAB-publicData-Hiseq-100x/outcome/GIAB-publicData-Hiseq-100x-HG006/.
    python .../03_make_pipeline_dirs.py --tier germline --dry-run                     # 2차 germline만 계획
    python .../03_make_pipeline_dirs.py --qsub --sampleinfo sampleinfo/sample_information_nbb2.xlsx   # 전부 + 병합 잡 + sampleinfo 행 추가
- **파이프라인이 이미 손댄 set dir(__DONE__·error_list.txt·tmp/·.snakemake/·analyze_meta/ 중 하나라도 있음)에 새 샘플 dir을 만들려 하면
  아무것도 만들지 않고 멈춘다.** 돌고 있거나 끝난 세트의 의미를 바꾸지 않기 위해서다. 이미 있는 샘플 dir의 재실행은 괜찮다.
- 병합은 샘플 dir의 `.concat.lock/`으로 한 번에 하나만 돈다(두 cat이 같은 .part에 쓰면 크기는 맞는데 내용이 섞인다). `--qsub`은 lock이 있거나
  `concat.jobid`의 잡이 아직 qstat에 있으면 다시 내지 않는다.
- `--sampleinfo <xlsx>`: 파이프라인이 읽는 샘플 정보 표(SET_ID/DNA_ID/SEX/note)에 계획된 샘플 중 없는 DNA_ID만 덧붙인다. 먼저
  `<xlsx>.bak-<시각>`으로 복사하고, 같은 DNA_ID가 다른 SET_ID·SEX로 이미 있으면 쓰지 않고 멈춘다. `--dry-run`이면 추가할 행만 보여준다.
"""
import argparse, csv, os, re, shutil, subprocess, sys, time
from collections import Counter, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
SHEETS = os.path.join(ROOT, "samplesheets")

# samplesheet dataset -> 디렉토리 이름 뒷부분 (2026-09-18 사용자 지정, 2026-09-19 `_`→`-`). 1차 22 샘플 전용.
# more_run_table.tsv에 있는 dsid(2차, HG002 NIST_BGIseq·Element_AVITI_20231018 포함)는 이 표보다 먼저 그 표의 set/label을 따른다.
DIRNAME = {
    "NovaSeq_PCRfree_30x":    "Novaseq6000-PCRfree-30x",
    "NovaSeqX_30x":           "NovaseqX-30x",
    "Illumina_PCRfree_30x":   "Hiseq-subsampled-30x",
    "HiSeq300x":              "Hiseq-300x",
    "Illumina_2x250":         "Illumina-250PE",
    "MGISEQ2000_PCRfree":     "MGISEQ2000-PCRfree",
    "BGISEQ500":              "BGISEQ500",
    "Element_AVITI_20240920": "Element-AVITI",
}
SEP = "-"   # 이름 구분자. `_`는 mate 접미에만
MORE_TABLE = os.path.join(ROOT, "more_run_table.tsv")   # 2차 set/label 배정 (make_more_samplesheets.py)
# 이 중 하나라도 set dir에 있으면 파이프라인이 그 set을 이미 돌렸거나 돌리는 중이다
PIPELINE_MARKERS = ("__DONE__", "error_list.txt", "tmp", ".snakemake", "analyze_meta")
# 1차 22 샘플(HG002~4)은 more_run_table에 없어 성별을 여기서 준다
SEX_FIRST = {"HG002": "male", "HG003": "male", "HG004": "female"}

CONCAT_SH = r"""#!/bin/bash
# {name}: FASTQ {n}쌍을 R1/R2 각각 하나로 cat 병합한다 (03_make_pipeline_dirs.py 생성, 2026-09-18 결정).
# 순서는 concat_R1.list / concat_R2.list 가 같다(unit, 경로 순). gzip 멀티멤버라 이어붙여도 유효하다.
# 재실행 안전: 출력 크기가 입력 합과 같으면 건너뛴다. 단독 실행: nohup bash concat.sh > concat.log 2>&1 &
set -euo pipefail
cd "$(dirname "$0")"
# 한 번에 하나만: 두 cat이 같은 .part에 쓰면 크기는 맞는데 내용이 섞여 크기 검증을 통과한다. mkdir은 Lustre에서도 원자적이다(flock은 마운트 옵션 따라 안 됨)
lock=".concat.lock"
if ! mkdir "$lock" 2>/dev/null; then
    echo "LOCKED: $(cat "$lock/owner" 2>/dev/null || echo '?') — 다른 병합이 도는 중. 그 잡이 죽었으면 rm -r $PWD/$lock 후 재실행"; exit 3
fi
echo "$(hostname) pid=$$ job=${{JOB_ID:-none}} $(date '+%F %T')" > "$lock/owner"
trap 'rm -rf "$lock"' EXIT
for m in 1 2; do
    out="{name}_${{m}}.fastq.gz"; list="concat_R${{m}}.list"
    want=$(xargs -d '\n' stat -c%s < "$list" | awk '{{s+=$1}} END{{print s}}')
    if [ -f "$out" ] && [ "$(stat -c%s "$out")" = "$want" ]; then echo "SKIP $out (크기 일치 $want)"; continue; fi
    echo "CAT  $(wc -l < "$list") files -> $out ($(awk -v b="$want" 'BEGIN{{printf "%.1f", b/1073741824}}') GiB) $(date '+%F %T')"
    xargs -d '\n' cat < "$list" > "$out.part"
    got=$(stat -c%s "$out.part")
    [ "$got" = "$want" ] || {{ echo "ERROR: 크기 불일치 $out (got=$got want=$want)"; exit 1; }}
    mv "$out.part" "$out"; echo "OK   $out $(date '+%F %T')"
done
# zcat | head -1 은 head가 먼저 끝나 zcat이 SIGPIPE(141)로 죽는다 — pipefail+errexit 아래서는 스크립트가 여기서 죽는다(2026-09-18 실측, 병합은 이미 끝난 뒤).
# (zcat || true) 로 감싸 SIGPIPE를 정상 종료로 만든다. HARVEST의 pipefail 교훈과 같은 함정.
echo "첫 리드 이름(짝이어야 함):"
( zcat "{name}_1.fastq.gz" || true ) | head -1
( zcat "{name}_2.fastq.gz" || true ) | head -1
echo "DONE {name} $(date '+%F %T')"
"""

QSUB_SH = """#!/bin/bash
#$ -N cat.{name}
#$ -q {queue}
#$ -S /bin/bash
#$ -j y
#$ -o {sdir}/concat.$JOB_ID.log
#$ -l h_vmem=4G
#$ -l h='{hosts}'
#$ -wd {sdir}
bash {sdir}/concat.sh
"""

ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
ap.add_argument("--dest", default=".", help="디렉토리를 만들 위치 (기본: 현재 디렉토리)")
ap.add_argument("--prefix", default="GIAB-publicData", help="디렉토리·샘플 이름 접두 (기본 GIAB-publicData; `_`는 mate 접미에만 쓰므로 접두에도 넣지 말 것)")
ap.add_argument("--only", action="append", default=[], help="이 데이터셋 디렉토리 이름(뒷부분)만. 반복 가능")
ap.add_argument("--hard", action="store_true", help="1쌍짜리를 심볼릭 링크 대신 하드링크로. 원본과 같은 파일시스템이어야 한다")
ap.add_argument("--run-concat", action="store_true", help="concat.sh 전부를 여기서 실행하고 끝날 때까지 기다린다 (--jobs 병렬)")
ap.add_argument("--jobs", type=int, default=3, help="--run-concat 동시 실행 수 (기본 3; 로그인 노드 I/O 고려)")
ap.add_argument("--qsub", action="store_true", help="concat.sh 를 샘플별 SGE 잡으로 제출 (이미 완료된 샘플은 건너뜀)")
# 쓸 수 있는 노드 집합은 고정이 아니다 — 2026-09-22에 shepherd 3대에서 octopus 2 + shepherd 2로 바뀌었고
# 큐도 둘로 갈렸다. phase1/2 의 env.sh 와 같은 값을 기본으로 두되, 환경변수로 덮는 쪽이 정상 경로다.
# 큐 x 노드가 안 겹치면 잡은 에러 없이 qw 로 영원히 남는다 — 제출 뒤 qstat 로 상태를 확인할 것.
ap.add_argument("--queue", default=os.environ.get("SGE_QUEUE", "shepherd.q,octopus.q"))
ap.add_argument("--hosts", default=os.environ.get("SGE_HOSTS", "(octopus-2-8|octopus-2-9|shepherd-1-8|shepherd-1-9)"))
ap.add_argument("--prune", action="store_true", help="관리하는 샘플 dir 안에서 계획에 없는 심볼릭 링크를 지운다(일반 파일·병합 결과는 안 건드림)")
ap.add_argument("--tier", choices=("all", "germline", "somatic"), default="all",
                help="germline = 정답셋 있는 HG001~7(1차 22 샘플 포함), somatic = HG008/HG009 (기본 all)")
ap.add_argument("--sampleinfo", metavar="XLSX", help="이 샘플 정보 xlsx에 없는 DNA_ID 행을 덧붙인다(백업 후). --dry-run이면 보여주기만")
ap.add_argument("--dry-run", action="store_true", help="만들지 않고 계획만 출력")
ap.add_argument("--sheets", default=SHEETS, help=argparse.SUPPRESS)   # 테스트용: samplesheet 디렉토리 바꿔치기
ap.add_argument("--more-table", default=MORE_TABLE, help=argparse.SUPPRESS)   # 테스트용
a = ap.parse_args()
SHEETS = a.sheets
if "_" in a.prefix:
    sys.exit(f"ERROR: 접두에 '_'가 있다({a.prefix}). '_'는 mate 접미(_1/_2.fastq.gz)에만 쓴다 — 예: GIAB-publicData")
BASH = os.environ.get("CONCAT_BASH", "bash")   # 테스트용: Windows에서 Git Bash 경로 지정
if a.run_concat and a.qsub:
    sys.exit("--run-concat 과 --qsub 은 하나만")


def read_sheet(path):
    rows = [l for l in open(path, encoding="utf-8") if not l.startswith("#")]
    return list(csv.DictReader(rows))


def concat_done(sdir, name, pairs):
    """출력 두 파일이 있고 크기가 입력 합과 같으면 True (concat.sh의 SKIP 조건과 동일)"""
    for m, idx in (("1", 0), ("2", 1)):
        out = os.path.join(sdir, f"{name}_{m}.fastq.gz")
        if not os.path.isfile(out):
            return False
        try:
            want = sum(os.path.getsize(p[idx]) for p in pairs)
        except OSError:
            return False
        if os.path.getsize(out) != want:
            return False
    return True


def read_more(path):
    """more_run_table.tsv -> {dsid: row}. set/label에 `_`·`.`가 있으면 멈춘다(파이프라인이 `_`로 자른다)"""
    if not os.path.exists(path):
        return {}
    out = {}
    for r in csv.DictReader(open(path, encoding="utf-8"), delimiter="\t"):
        for k in ("set", "label"):
            if "_" in r[k] or "." in r[k]:
                sys.exit(f"ERROR: {path}: {r['dsid']} 의 {k}='{r[k]}' 에 '_' 또는 '.'")
        out[r["dsid"]] = r
    return out


def is_touched(set_dir):
    """파이프라인이 이 set을 이미 돌렸거나 돌리는 중인가"""
    return any(os.path.lexists(os.path.join(set_dir, m)) for m in PIPELINE_MARKERS)


MORE = read_more(a.more_table)
plan = []      # (link_path, target)          — 1쌍짜리
concat = []    # (sample_dir, name, [(r1, r2), ...] 고정 순서) — 여러 쌍
missing = []
samples = []   # (set dir 이름, 샘플 이름, 성별, note) — sampleinfo 행
for fn in sorted(os.listdir(SHEETS)):
    if not fn.endswith(".csv"):
        continue
    dsid = fn[:-4]
    sample, dataset = dsid.split(".", 1)
    if dsid in MORE:
        m = MORE[dsid]
        dname, tier, sex, note = f"{a.prefix}{SEP}{m['set']}", m["tier"], m["sex"], m["label"]
        short, name = m["set"], f"{a.prefix}{SEP}{m['set']}{SEP}{m['label']}"
    elif dataset in DIRNAME:
        dname, tier, sex, note = f"{a.prefix}{SEP}{DIRNAME[dataset]}", "germline", SEX_FIRST.get(sample, ""), sample
        short, name = DIRNAME[dataset], f"{a.prefix}{SEP}{DIRNAME[dataset]}{SEP}{sample}"
    else:
        continue
    if a.tier != "all" and tier != a.tier:
        continue
    if a.only and short not in a.only and dname not in a.only:
        continue
    rows = read_sheet(os.path.join(SHEETS, fn))
    sdir = os.path.join(a.dest, dname, "outcome", name)
    samples.append((dname, name, sex, note, sdir))
    if len(rows) > 1:
        pairs = [(r["fastq_1"], r["fastq_2"]) for r in sorted(rows, key=lambda r: (r["unit"], r["fastq_1"]))]
        concat.append((sdir, name, pairs))
        if not a.dry_run:
            missing += [f for pr in pairs for f in pr if not os.path.exists(f)]
        continue
    r = rows[0]
    for mate in ("1", "2"):
        src = r[f"fastq_{mate}"]
        plan.append((os.path.join(sdir, f"{name}_{mate}.fastq.gz"), src))
        if not a.dry_run and not os.path.exists(src):
            missing.append(src)

dups = [l for l, c in Counter(l for l, _ in plan).items() if c > 1]
dups += [d for d, c in Counter(s[4] for s in samples).items() if c > 1]
if dups:
    print(f"ERROR: 링크·샘플 dir 이름이 겹친다 {len(dups)}개. 처음 5개:", *dups[:5], sep="\n  "); sys.exit(1)
if missing:
    print(f"ERROR: 원본 FASTQ {len(missing)}개가 없다 (DATA_ROOT 확인). 처음 5개:", *missing[:5], sep="\n  ")
    sys.exit(1)
# 새 샘플 dir을 파이프라인이 이미 손댄 set에 만들려는가 — dry-run에서도 알린다(만들기 전에 멈추는 것이 목적)
intrude = [s[4] for s in samples if not os.path.isdir(s[4]) and is_touched(os.path.join(a.dest, s[0]))]
if intrude:
    print(f"ERROR: 파이프라인이 이미 돌았거나 도는 set에 새 샘플 {len(intrude)}개를 넣으려 한다 — 아무것도 만들지 않았다.",
          "more_run_table.tsv에서 새 set 이름을 줄 것. 처음 5개:", *intrude[:5], sep="\n  ")
    sys.exit(1)


def update_sampleinfo(xlsx, samples, dry):
    """xlsx의 SET_ID/DNA_ID 표에 없는 DNA_ID만 덧붙인다. 충돌(같은 DNA_ID, 다른 SET_ID·SEX)이면 쓰지 않고 멈춘다"""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]]
        if "SET_ID" in hdr and "DNA_ID" in hdr:
            break
    else:
        sys.exit(f"ERROR: {xlsx}에 SET_ID/DNA_ID 머리행이 있는 시트가 없다")
    col = {h: i for i, h in enumerate(hdr)}
    have = {r[col["DNA_ID"]]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[col["DNA_ID"]]}
    add, bad = [], []
    for dname, name, sex, note, _ in samples:
        if name in have:
            e = have[name]
            if e[col["SET_ID"]] != dname or ("SEX" in col and str(e[col["SEX"]] or "").lower() != sex):
                bad.append(f"{name}: 표에는 SET_ID={e[col['SET_ID']]} SEX={e[col['SEX']] if 'SEX' in col else '?'}, 계획은 {dname} {sex}")
        elif not sex:
            bad.append(f"{name}: 성별을 모른다 — more_run_table.tsv에 sex를 채울 것")
        else:
            add.append({"SET_ID": dname, "DNA_ID": name, "SEX": sex, "note": note})
    if bad:
        print("ERROR: sampleinfo 충돌 — xlsx를 건드리지 않았다:", *bad[:10], sep="\n  "); sys.exit(1)
    print(f"sampleinfo {xlsx}: 이미 있음 {len(samples) - len(add)}, 추가 {len(add)}{' (dry-run, 안 씀)' if dry else ''}")
    for r in add[:60]:
        print(f"  + {r['SET_ID']}\t{r['DNA_ID']}\t{r['SEX']}\t{r['note']}")
    if dry or not add:
        return
    bak = f"{xlsx}.bak-{time.strftime('%Y%m%d-%H%M%S')}"
    shutil.copy2(xlsx, bak)
    for r in add:
        ws.append([r.get(h, "") for h in hdr])
    wb.save(xlsx)
    print(f"  저장함 (백업 {bak})")


if a.sampleinfo:
    update_sampleinfo(a.sampleinfo, samples, a.dry_run)

# ---- 1쌍짜리: 링크 ------------------------------------------------------------------
made = skipped = pruned = 0
planned = {os.path.abspath(l) for l, _ in plan}
managed = {os.path.dirname(l) for l, _ in plan} | {c[0] for c in concat}
if a.prune and not a.dry_run:
    for d in managed:
        if not os.path.isdir(d):
            continue
        for fn in os.listdir(d):
            f = os.path.join(d, fn)
            if os.path.islink(f) and os.path.abspath(f) not in planned:
                os.remove(f); pruned += 1; print(f"PRUNE {os.path.relpath(f, a.dest)}")
for link, src in plan:
    if a.dry_run:
        continue
    os.makedirs(os.path.dirname(link), exist_ok=True)
    if os.path.lexists(link):
        cur = os.readlink(link) if os.path.islink(link) else (link if a.hard and os.path.samefile(link, src) else None)
        if cur == src or (a.hard and cur == link):
            skipped += 1; continue
        sys.exit(f"ERROR: 이미 있는데 다른 대상을 가리킨다: {link} -> {cur if cur else '(하드링크/일반 파일, 원본과 다름)'} (원한 것: {src})")
    if a.hard:
        os.link(src, link)
    else:
        os.symlink(src, link)
    made += 1

# ---- 여러 쌍: concat.sh 작성 ------------------------------------------------------------
for sdir, name, pairs in concat:
    if a.dry_run:
        continue
    os.makedirs(sdir, exist_ok=True)
    with open(os.path.join(sdir, "concat_R1.list"), "w", newline="\n") as f1, \
         open(os.path.join(sdir, "concat_R2.list"), "w", newline="\n") as f2:
        for r1, r2 in pairs:
            f1.write(r1 + "\n"); f2.write(r2 + "\n")
    sh = os.path.join(sdir, "concat.sh")
    with open(sh, "w", newline="\n") as f:
        f.write(CONCAT_SH.format(name=name, n=len(pairs)))
    os.chmod(sh, 0o755)

n_ds = len({os.path.dirname(os.path.dirname(d)) for d in managed})
print(f"{'DRY-RUN ' if a.dry_run else ''}링크 {len(plan)}개({made} 생성, {skipped} 기존, {pruned} 정리) · 병합 {len(concat)} 샘플 · "
      f"데이터셋 dir {n_ds}개 under {os.path.abspath(a.dest)}")
for d in sorted({os.path.dirname(l) for l, _ in plan}):
    print(f"  LINK   {os.path.relpath(d, a.dest)}")
for sdir, name, pairs in concat:
    gib = sum(os.path.getsize(p) for pr in pairs for p in pr) / 2**30 if not a.dry_run else 0
    state = "" if a.dry_run else (" — 이미 완료" if concat_done(sdir, name, pairs) else f" — {gib:.0f} GiB, concat.sh 작성됨")
    print(f"  CONCAT {os.path.relpath(sdir, a.dest)}  ({len(pairs)} pairs -> {name}_1/2.fastq.gz{state})")

if a.dry_run or not concat:
    sys.exit(0)
todo = [(sdir, name) for sdir, name, pairs in concat if not concat_done(sdir, name, pairs)]

# ---- 병합 실행 -------------------------------------------------------------------------
def job_alive(sdir):
    """concat.jobid 의 잡이 아직 qstat 에 있으면 그 번호"""
    p = os.path.join(sdir, "concat.jobid")
    if not os.path.exists(p):
        return None
    jid = open(p).read().strip()
    if not jid.isdigit():
        return None
    r = subprocess.run(["qstat", "-j", jid], capture_output=True, text=True, stdin=subprocess.DEVNULL)
    return jid if r.returncode == 0 else None


if a.qsub:
    n_sub = 0
    for sdir, name in todo:
        if os.path.isdir(os.path.join(sdir, ".concat.lock")):
            print(f"SKIP {name}: .concat.lock 있음(병합 중이거나 죽은 잡의 흔적 — owner: "
                  f"{open(os.path.join(sdir, '.concat.lock', 'owner')).read().strip() if os.path.exists(os.path.join(sdir, '.concat.lock', 'owner')) else '?'})")
            continue
        try:
            alive = job_alive(sdir)
        except FileNotFoundError:
            sys.exit("ERROR: qstat/qsub 을 찾을 수 없다 (로그인 노드에서 실행할 것)")
        if alive:
            print(f"SKIP {name}: 잡 {alive} 가 아직 큐에 있다"); continue
        job = os.path.join(sdir, "concat.qsub.sh")
        with open(job, "w", newline="\n") as f:
            f.write(QSUB_SH.format(name=name, sdir=os.path.abspath(sdir), queue=a.queue, hosts=a.hosts))
        out = subprocess.run(["qsub", job], capture_output=True, text=True, stdin=subprocess.DEVNULL)
        msg = (out.stdout or out.stderr).strip()
        m = re.search(r"Your job (\d+)", msg)
        if out.returncode != 0 or not m:
            print(f"FAIL qsub {name}: {msg}"); continue
        with open(os.path.join(sdir, "concat.jobid"), "w") as f:
            f.write(m.group(1) + "\n")
        n_sub += 1
        print(f"QSUB {name}: {msg}")
    print(f"{n_sub}/{len(todo)} 잡 제출. 상태: qstat  · 결과: <샘플 dir>/concat.<JOB_ID>.log 마지막 줄이 'DONE'")
elif a.run_concat:
    running = {}   # name -> (Popen, sdir, log)
    failed = []
    todo = list(todo)
    print(f"병합 {len(todo)} 샘플, 동시 {a.jobs}개. 로그: <샘플 dir>/concat.log")
    while todo or running:
        while todo and len(running) < a.jobs:
            sdir, name = todo.pop(0)
            log = open(os.path.join(sdir, "concat.log"), "a")
            pr = subprocess.Popen([BASH, os.path.join(os.path.abspath(sdir), "concat.sh").replace(os.sep, "/")], stdout=log, stderr=subprocess.STDOUT)
            running[name] = (pr, sdir, log)
            print(f"START {name} {time.strftime('%F %T')}", flush=True)
        for name in list(running):
            pr, sdir, log = running[name]
            if pr.poll() is None:
                continue
            log.close(); del running[name]
            print(f"{'OK   ' if pr.returncode == 0 else 'FAIL '}{name} rc={pr.returncode} {time.strftime('%F %T')}", flush=True)
            if pr.returncode != 0:
                failed.append(name)
        time.sleep(5)
    done = sum(1 for sdir, name, pairs in concat if concat_done(sdir, name, pairs))
    print(f"병합 끝: 완료 {done}/{len(concat)} 샘플(출력 크기 = 입력 합), 이번 실행 실패 {len(failed)}개"
          + (f" — {', '.join(failed)} (concat.log 확인, 재실행하면 이어서 함)" if failed else ""))
    sys.exit(1 if failed or done < len(concat) else 0)
else:
    done = len(concat) - len(todo)
    print(f"병합 완료 {done}/{len(concat)} 샘플(출력 크기 = 입력 합).")
    print(f"병합 미실행 {len(todo)} 샘플. 한 번에 전부: --run-concat (--jobs N) 또는 --qsub. 개별: nohup bash <샘플 dir>/concat.sh > <샘플 dir>/concat.log 2>&1 &")
